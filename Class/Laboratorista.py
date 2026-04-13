"""Funciones de apoyo para el módulo de laboratorista.

Aquí se concentra la consulta de equipos, préstamos y la generación de la
bitácora en Excel para el laboratorio.
"""

from db.Conection import obtener_cursor
from openpyxl import load_workbook
from datetime import datetime

class LaboratoristaManager:
    """Agrupa consultas y operaciones del módulo de préstamo de equipos."""
    def obtener_equipos(self):
        conexion, cursor = obtener_cursor()
        cursor.execute("""
            SELECT 
                e.id,
                e.codigo_equipo,
                e.nombre_equipo,
                e.categoria,
                e.marca,
                e.modelo,
                e.numero_serie,
                e.descripcion,
                e.estado,
                e.ubicacion,
                e.fecha_registro,
                COALESCE(a.nombre_completo, ex.nombre_completo, pe.nombre_completo) AS nombre_usuario
            FROM equipos_laboratorio e
            LEFT JOIN prestamos_equipo p 
                ON p.equipo_id = e.id AND p.estado = 'activo'
            LEFT JOIN alumnos a 
                ON a.matricula = p.matricula_usuario AND p.tipo_usuario = 'alumno'
            LEFT JOIN externos ex 
                ON ex.matricula = p.matricula_usuario AND p.tipo_usuario = 'externo'
            LEFT JOIN personal pe 
                ON pe.matricula = p.matricula_usuario AND p.tipo_usuario IN ('tutor', 'laboratorista')
            ORDER BY e.nombre_equipo ASC
        """)
        datos = cursor.fetchall()
        cursor.close()
        conexion.close()
        return datos

    def obtener_equipos_disponibles(self):
        conexion, cursor = obtener_cursor()
        cursor.execute("""
            SELECT id, codigo_equipo, nombre_equipo, categoria, marca, modelo, numero_serie, descripcion, estado, ubicacion
            FROM equipos_laboratorio
            WHERE estado = 'disponible'
            ORDER BY nombre_equipo ASC
        """)
        datos = cursor.fetchall()
        cursor.close()
        conexion.close()
        return datos

    def obtener_prestamos_activos(self):
        conexion, cursor = obtener_cursor()
        cursor.execute("""
            SELECT 
                p.id,
                p.equipo_id,
                p.matricula_usuario,
                p.tipo_usuario,
                p.fecha_prestamo,
                p.fecha_limite,
                p.estado,
                p.observaciones,
                e.codigo_equipo,
                e.nombre_equipo,
                e.categoria,
                COALESCE(a.nombre_completo, ex.nombre_completo, pe.nombre_completo) AS nombre_usuario,
                c.nombre AS carrera
            FROM prestamos_equipo p
            INNER JOIN equipos_laboratorio e ON e.id = p.equipo_id
            LEFT JOIN alumnos a ON a.matricula = p.matricula_usuario AND p.tipo_usuario = 'alumno'
            LEFT JOIN carreras c ON c.id = a.carrera_id
            LEFT JOIN externos ex ON ex.matricula = p.matricula_usuario AND p.tipo_usuario = 'externo'
            LEFT JOIN personal pe ON pe.matricula = p.matricula_usuario AND p.tipo_usuario IN ('tutor', 'laboratorista')
            WHERE p.estado = 'activo'
            ORDER BY p.fecha_prestamo DESC
        """)
        datos = cursor.fetchall()
        cursor.close()
        conexion.close()
        return datos

    def registrar_equipo(self, datos):
        try:
            conexion, cursor = obtener_cursor()
            cursor.execute("""
                INSERT INTO equipos_laboratorio
                (codigo_equipo, nombre_equipo, categoria, marca, modelo, numero_serie, descripcion, estado, ubicacion)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
            """, (
                datos["codigo_equipo"],
                datos["nombre_equipo"],
                datos["categoria"],
                datos["marca"],
                datos["modelo"],
                datos["numero_serie"],
                datos["descripcion"],
                datos["estado"],
                datos["ubicacion"]
            ))
            conexion.commit()
            cursor.close()
            conexion.close()
            return {"status": "ok"}
        except Exception as e:
            try:
                cursor.close()
                conexion.close()
            except:
                pass

            if "Duplicate entry" in str(e):
                return {"status": "exists"}

            return {"status": "error"}

    def buscar_usuario_por_matricula(self, matricula):
        conexion, cursor = obtener_cursor()

        cursor.execute("""
            SELECT 
                a.matricula,
                a.nombre_completo,
                a.email,
                'alumno' AS tipo_usuario,
                c.nombre AS carrera
            FROM alumnos a
            INNER JOIN carreras c ON c.id = a.carrera_id
            WHERE a.matricula = %s AND a.activo = 1
        """, (matricula,))
        alumno = cursor.fetchone()

        if alumno:
            cursor.close()
            conexion.close()
            return alumno

        cursor.execute("""
            SELECT 
                matricula,
                nombre_completo,
                email,
                'externo' AS tipo_usuario,
                'Externo' AS carrera
            FROM externos
            WHERE matricula = %s AND activo = 1
        """, (matricula,))
        externo = cursor.fetchone()

        if externo:
            cursor.close()
            conexion.close()
            return externo

        cursor.execute("""
            SELECT 
                matricula,
                nombre_completo,
                email,
                tipo AS tipo_usuario,
                'Personal' AS carrera
            FROM personal
            WHERE matricula = %s AND activo = 1
        """, (matricula,))
        personal = cursor.fetchone()

        cursor.close()
        conexion.close()
        return personal

    def registrar_prestamo(self, equipo_id, matricula_usuario, tipo_usuario, fecha_limite, observaciones):
        try:
            conexion, cursor = obtener_cursor()

            cursor.execute("""
                SELECT id, estado
                FROM equipos_laboratorio
                WHERE id = %s
            """, (equipo_id,))
            equipo = cursor.fetchone()

            if not equipo:
                cursor.close()
                conexion.close()
                return {"status": "team_not_found"}

            if equipo["estado"] != "disponible":
                cursor.close()
                conexion.close()
                return {"status": "not_available"}

            cursor.execute("""
                SELECT id
                FROM prestamos_equipo
                WHERE equipo_id = %s AND estado = 'activo'
            """, (equipo_id,))
            prestamo_activo = cursor.fetchone()

            if prestamo_activo:
                cursor.close()
                conexion.close()
                return {"status": "already_loaned"}

            cursor.execute("""
                INSERT INTO prestamos_equipo
                (equipo_id, matricula_usuario, tipo_usuario, fecha_limite, estado, observaciones)
                VALUES (%s, %s, %s, %s, 'activo', %s)
            """, (
                equipo_id,
                matricula_usuario,
                tipo_usuario,
                fecha_limite if fecha_limite != "" else None,
                observaciones if observaciones != "" else None
            ))

            cursor.execute("""
                UPDATE equipos_laboratorio
                SET estado = 'prestado'
                WHERE id = %s
            """, (equipo_id,))

            conexion.commit()
            cursor.close()
            conexion.close()
            return {"status": "ok"}

        except Exception:
            try:
                cursor.close()
                conexion.close()
            except:
                pass
            return {"status": "error"}

    def obtener_historial_bitacora(self):
        conexion, cursor = obtener_cursor()
        cursor.execute("""
            SELECT
                p.fecha_prestamo,
                p.fecha_limite,
                p.fecha_devolucion,
                p.estado,
                p.observaciones,
                p.tipo_usuario,
                p.matricula_usuario,
                e.codigo_equipo,
                e.nombre_equipo,
                e.categoria,
                COALESCE(a.nombre_completo, ex.nombre_completo, pe.nombre_completo) AS nombre_usuario,
                COALESCE(c.nombre, 'Externo') AS carrera
            FROM prestamos_equipo p
            INNER JOIN equipos_laboratorio e ON e.id = p.equipo_id
            LEFT JOIN alumnos a ON a.matricula = p.matricula_usuario AND p.tipo_usuario = 'alumno'
            LEFT JOIN carreras c ON c.id = a.carrera_id
            LEFT JOIN externos ex ON ex.matricula = p.matricula_usuario AND p.tipo_usuario = 'externo'
            LEFT JOIN personal pe ON pe.matricula = p.matricula_usuario AND p.tipo_usuario IN ('tutor', 'laboratorista')
            ORDER BY p.fecha_prestamo DESC
        """)
        datos = cursor.fetchall()
        cursor.close()
        conexion.close()
        return datos

    def generar_excel_bitacora(self, nombre_laboratorio, periodo, responsable, template_path, output_path):
        registros = self.obtener_historial_bitacora()

        wb = load_workbook(template_path)
        ws = wb.active

        ws["B5"] = f"LABORATORIO: {nombre_laboratorio}"
        ws["K5"] = f"PERIODO: {periodo}"

        fila = 9

        for item in registros:
            fecha = item["fecha_prestamo"].strftime("%d/%m/%Y") if item["fecha_prestamo"] else ""
            equipo = f'{item["codigo_equipo"]} - {item["nombre_equipo"]}'
            actividad = item["observaciones"] if item["observaciones"] else "Préstamo de equipo"

            ws[f"B{fila}"] = fecha
            ws[f"C{fila}"] = responsable
            ws[f"D{fila}"] = item["carrera"] if item["carrera"] else ""
            ws[f"E{fila}"] = item["tipo_usuario"].capitalize()
            ws[f"F{fila}"] = "Préstamo de laboratorio"
            ws[f"G{fila}"] = actividad
            ws[f"H{fila}"] = 1
            ws[f"I{fila}"] = 0
            ws[f"J{fila}"] = 0
            ws[f"K{fila}"] = item["nombre_usuario"] if item["nombre_usuario"] else item["matricula_usuario"]
            ws[f"L{fila}"] = ""
            ws[f"M{fila}"] = equipo

            fila += 1

            if fila > 35:
                break

        wb.save(output_path)
        return output_path


    def actualizar_estado_prestamos(self):
        conexion, cursor = obtener_cursor()

        cursor.execute("""
            SELECT p.id, p.equipo_id, p.fecha_limite, p.estado, e.nombre_equipo
            FROM prestamos_equipo p
            INNER JOIN equipos_laboratorio e ON e.id = p.equipo_id
            WHERE p.estado = 'activo' AND p.fecha_limite IS NOT NULL
        """)
        prestamos = cursor.fetchall()

        ahora = datetime.now()
        vencidos = []

        for prestamo in prestamos:
            if prestamo["fecha_limite"] and ahora > prestamo["fecha_limite"]:
                cursor.execute("""
                    UPDATE prestamos_equipo
                    SET estado = 'vencido'
                    WHERE id = %s
                """, (prestamo["id"],))

                cursor.execute("""
                    UPDATE equipos_laboratorio
                    SET estado = 'disponible'
                    WHERE id = %s
                """, (prestamo["equipo_id"],))

                vencidos.append(prestamo["nombre_equipo"])

        conexion.commit()
        cursor.close()
        conexion.close()
        return vencidos
