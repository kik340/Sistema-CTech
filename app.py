"""
Archivo principal de Flask para el proyecto QR ITSA.

Este archivo centraliza la configuración de la aplicación, el control de sesiones
y el enrutamiento por tipo de usuario. La lógica original se conserva, pero las
rutas se agrupan por módulos funcionales para que sea más fácil ubicar el flujo
de login, alumno, tutor y laboratorista.
"""

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session,
    send_file,
    jsonify,
)
import io
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

import qrcode

from Class.Laboratorista import LaboratoristaManager
from Class.PDFManager import PdfQrAlumno, PdfSeguimientoAcademico
from Class.Tutor import TutorManager
from Class.User import User

# -------------------------------------------------------------------
# Configuración base de la aplicación y carga de gestores principales
# -------------------------------------------------------------------

app = Flask(__name__)
app.secret_key = "clave_secreta_123"
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SECURE"] = False

usuario = User()
tutor_manager = TutorManager()
laboratorista_manager = LaboratoristaManager()

UPLOAD_FOLDER = "static/uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)


# -------------------------------------------------------------------
# Funciones auxiliares compartidas entre los distintos módulos
# -------------------------------------------------------------------

def login_requerido(tipo=None):
    """Valida si existe sesión activa y, opcionalmente, si coincide con el tipo de usuario esperado."""
    if "tipo" not in session:
        return False
    if tipo and session["tipo"] != tipo:
        return False
    return True


def obtener_foto_usuario(matricula):
    """Busca la imagen del usuario en la carpeta de uploads y devuelve su ruta pública con versión para evitar caché."""
    carpeta = app.config["UPLOAD_FOLDER"]
    extensiones = ["png", "jpg", "jpeg", "webp"]

    for ext in extensiones:
        ruta = os.path.join(carpeta, f"{matricula}.{ext}")
        if os.path.exists(ruta):
            version = int(os.path.getmtime(ruta))
            return url_for("static", filename=f"uploads/{matricula}.{ext}", v=version)

    return url_for("static", filename="uploads/none.png")

def texto_si_no(valor):
    """Convierte valores binarios del seguimiento a textos legibles para exportación."""
    return "Sí" if int(valor or 0) == 1 else "No"


def construir_excel_seguimiento(seguimiento, filas):
    """Arma un archivo Excel en memoria con el contenido completo del seguimiento académico."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Seguimiento"

    rojo = "B30000"
    blanco = "FFFFFF"
    gris = "F4F4F4"
    borde = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws.merge_cells("A1:U1")
    ws["A1"] = "Seguimiento académico"
    ws["A1"].font = Font(bold=True, size=16, color=blanco)
    ws["A1"].fill = PatternFill("solid", fgColor=rojo)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    info = [
        ("Grupo ID", seguimiento.get("grupo_id", "")),
        ("Parcial", seguimiento.get("parcial", "")),
        ("Total estudiantes", seguimiento.get("total_estudiantes", "")),
        ("Observaciones generales", seguimiento.get("observaciones") or "Sin observaciones"),
    ]

    fila_inicio = 3
    for i, (label, value) in enumerate(info, start=fila_inicio):
        ws[f"A{i}"] = label
        ws[f"A{i}"].font = Font(bold=True)
        ws[f"A{i}"].fill = PatternFill("solid", fgColor=gris)
        ws[f"B{i}"] = value

    encabezados = [
        "Matrícula", "Nombre", "Asistencia", "Canalización", "Prob. económicos",
        "Prob. salud", "Prob. emocionales", "Prob. familiares", "Prob. académicos",
        "Beca JCF", "Beca manutención", "Beca JMF", "Otra beca", "Apoyo psicológico",
        "Apoyo médico", "Otro apoyo", "Promedio", "Reprobadas", "Aprobado",
        "Posible deserción", "Comentario"
    ]

    header_row = 8
    for col, titulo in enumerate(encabezados, start=1):
        cell = ws.cell(row=header_row, column=col, value=titulo)
        cell.font = Font(bold=True, color=blanco)
        cell.fill = PatternFill("solid", fgColor=rojo)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = borde

    for row_idx, fila in enumerate(filas, start=header_row + 1):
        valores = [
            fila.get("matricula", ""),
            fila.get("nombre_completo", ""),
            texto_si_no(fila.get("asistencia")),
            texto_si_no(fila.get("canalizacion_tutoria_individual")),
            texto_si_no(fila.get("problemas_economicos")),
            texto_si_no(fila.get("problemas_salud")),
            texto_si_no(fila.get("problemas_emocionales")),
            texto_si_no(fila.get("problemas_familiares")),
            texto_si_no(fila.get("problemas_academicos")),
            texto_si_no(fila.get("beca_jcf")),
            texto_si_no(fila.get("beca_manutencion")),
            texto_si_no(fila.get("beca_jmf")),
            texto_si_no(fila.get("beca_otro")),
            texto_si_no(fila.get("apoyo_psicologico")),
            texto_si_no(fila.get("apoyo_medico")),
            texto_si_no(fila.get("apoyo_otro")),
            fila.get("promedio", ""),
            fila.get("materias_reprobadas", 0),
            texto_si_no(fila.get("estatus_aprobado")),
            texto_si_no(fila.get("posible_desercion")),
            fila.get("comentario", ""),
        ]

        for col, valor in enumerate(valores, start=1):
            cell = ws.cell(row=row_idx, column=col, value=valor)
            cell.border = borde
            cell.alignment = Alignment(horizontal="center" if 3 <= col <= 20 else "left", vertical="center", wrap_text=True)

    widths = {
        "A": 16, "B": 28, "C": 12, "D": 14, "E": 16, "F": 14, "G": 16, "H": 16,
        "I": 16, "J": 12, "K": 16, "L": 12, "M": 12, "N": 16, "O": 14, "P": 12,
        "Q": 12, "R": 12, "S": 12, "T": 16, "U": 28
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A9"
    ws.sheet_view.showGridLines = False

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
def construir_datos_exportacion_desde_formulario():
    matriculas = request.form.getlist("alumno_matricula")

    alumnos = []
    detalle_map = {}

    for matricula in matriculas:
        nombre = request.form.get(f"nombre_{matricula}", "").strip()

        if not nombre:
            nombre = matricula

        alumnos.append({
            "matricula": matricula,
            "nombre_completo": nombre
        })

        detalle_map[matricula] = {
            "asistencia": 1 if request.form.get(f"asistencia_{matricula}") else 0,
            "canalizacion_tutoria_individual": 1 if request.form.get(f"canalizacion_{matricula}") else 0,
            "problemas_economicos": 1 if request.form.get(f"prob_economicos_{matricula}") else 0,
            "problemas_salud": 1 if request.form.get(f"prob_salud_{matricula}") else 0,
            "problemas_emocionales": 1 if request.form.get(f"prob_emocionales_{matricula}") else 0,
            "problemas_familiares": 1 if request.form.get(f"prob_familiares_{matricula}") else 0,
            "problemas_academicos": 1 if request.form.get(f"prob_academicos_{matricula}") else 0,
            "beca_jcf": 1 if request.form.get(f"beca_jcf_{matricula}") else 0,
            "beca_manutencion": 1 if request.form.get(f"beca_manutencion_{matricula}") else 0,
            "beca_jmf": 1 if request.form.get(f"beca_jmf_{matricula}") else 0,
            "beca_otro": 1 if request.form.get(f"beca_otro_{matricula}") else 0,
            "apoyo_psicologico": 1 if request.form.get(f"apoyo_psicologico_{matricula}") else 0,
            "apoyo_medico": 1 if request.form.get(f"apoyo_medico_{matricula}") else 0,
            "apoyo_otro": 1 if request.form.get(f"apoyo_otro_{matricula}") else 0,
            "promedio": request.form.get(f"promedio_{matricula}", "").strip(),
            "materias_reprobadas": request.form.get(f"reprobadas_{matricula}", "0").strip(),
            "estatus_aprobado": 1 if request.form.get(f"aprobado_{matricula}") else 0,
            "posible_desercion": 1 if request.form.get(f"desercion_{matricula}") else 0,
            "comentario": request.form.get(f"comentario_{matricula}", "").strip()
        }

    seguimiento = {
        "id": request.form.get("seguimiento_id"),
        "grupo_id": request.form.get("grupo_id", ""),
        "parcial": request.form.get("parcial", ""),
        "observaciones": request.form.get("observaciones", "").strip(),
        "total_estudiantes": len(alumnos)
    }

    return seguimiento, alumnos, detalle_map


# -------------------------------------------------------------------
# Acceso y registro público
# -------------------------------------------------------------------

@app.route("/")
def index():
    """Muestra la pantalla principal del sistema y limpia cualquier sesión previa antes de cargar las carreras activas."""
    session.clear()
    carreras = tutor_manager.obtener_carreras_activas()
    return render_template("index.html", carreras=carreras)

@app.route("/verificar_usuario", methods=["POST"])
def verificar_usuario():
    """Procesa el inicio de sesión, guarda los datos básicos en sesión y redirige al menú correspondiente según el tipo."""
    matricula = request.form.get("matricula", "").strip()
    password = request.form.get("password", "").strip()

    if matricula == "" or password == "":
        return redirect(url_for("index", r="empty_login"))

    resultado = usuario.login_por_matricula(matricula, password)

    if resultado["status"] != "ok":
        return redirect(url_for("index", r=resultado["status"]))

    session["matricula"] = resultado["matricula"]
    session["nombre"] = resultado["nombre"]
    session["email"] = resultado["email"]
    session["tipo"] = resultado["tipo"]
    session["carrera"] = resultado.get("carrera")

    if resultado["tipo"] == "alumno":
        return redirect(url_for("menu_alumno"))

    if resultado["tipo"] == "externo":
        return redirect(url_for("menu_alumno"))

    if resultado["tipo"] == "tutor":
        return redirect(url_for("menu_tutor"))

    if resultado["tipo"] == "laboratorista":
        return redirect(url_for("menu_laboratorista"))

    session.clear()
    return redirect(url_for("index", r="tipo_wrong"))

@app.route("/solicitar_registro", methods=["POST"])
def solicitar_registro():
    """Recibe la solicitud de registro pública, valida campos mínimos y delega la creación de la solicitud al gestor de usuarios."""
    nombre = request.form.get("nombre", "").strip()
    ap_paterno = request.form.get("ap_paterno", "").strip()
    ap_materno = request.form.get("ap_materno", "").strip()
    correo_temporal = request.form.get("correo_temporal", "").strip()
    password = request.form.get("password", "").strip()
    tipo_solicitud = request.form.get("tipo_solicitud", "").strip()
    carrera_id = request.form.get("carrera_id", "").strip()
    anio_ingreso = request.form.get("anio_ingreso", "").strip()
    motivo = request.form.get("motivo", "").strip()

    if (
        nombre == ""
        or ap_paterno == ""
        or ap_materno == ""
        or correo_temporal == ""
        or password == ""
        or tipo_solicitud == ""
    ):
        return redirect(url_for("index", r="empty_request"))

    if tipo_solicitud == "alumno" and (carrera_id == "" or anio_ingreso == ""):
        return redirect(url_for("index", r="empty_request"))

    if tipo_solicitud == "externo" and motivo == "":
        return redirect(url_for("index", r="empty_request"))

    try:
        anio_ingreso_int = int(anio_ingreso) if anio_ingreso != "" else None
    except:
        return redirect(url_for("index", r="empty_request"))

    datos = {
        "nombre": nombre,
        "ap_paterno": ap_paterno,
        "ap_materno": ap_materno,
        "correo_temporal": correo_temporal,
        "password": password,
        "tipo_solicitud": tipo_solicitud,
        "carrera_id": int(carrera_id) if carrera_id != "" else None,
        "anio_ingreso": anio_ingreso_int,
        "motivo": motivo if motivo != "" else None,
    }

    resultado = usuario.crear_solicitud(datos)
    return redirect(url_for("index", r=resultado["status"]))

@app.route("/cerrar_sesion")
def cerrar_sesion():
    """Limpia la sesión actual y regresa a la página principal."""
    session.clear()
    return redirect(url_for("index"))


# -------------------------------------------------------------------
# Operaciones compartidas de usuario autenticado
# -------------------------------------------------------------------

@app.route("/modificar_mi_usuario", methods=["POST"])
def modificar_mi_usuario():
    """Permite cambiar la contraseña del usuario autenticado y fuerza un nuevo inicio de sesión cuando el cambio es correcto."""
    if not login_requerido():
        return redirect(url_for("index"))

    nueva_password = request.form.get("nueva_password", "").strip()
    confirmar_password = request.form.get("confirmar_password", "").strip()

    if nueva_password == "" or confirmar_password == "":
        if session["tipo"] == "tutor":
            return redirect(url_for("menu_tutor", r="empty_request"))
        if session["tipo"] == "laboratorista":
            return redirect(url_for("menu_laboratorista", r="empty_request"))
        return redirect(url_for("menu_alumno", r="empty_request"))

    if nueva_password != confirmar_password:
        if session["tipo"] == "tutor":
            return redirect(url_for("menu_tutor", r="password_no_match"))
        if session["tipo"] == "laboratorista":
            return redirect(url_for("menu_laboratorista", r="password_no_match"))
        return redirect(url_for("menu_alumno", r="password_no_match"))

    resultado = usuario.cambiar_password_usuario(
        session["matricula"],
        session["tipo"],
        nueva_password
    )

    if resultado["status"] == "ok":
        session.clear()
        return redirect(url_for("index", r="password_changed"))

    if session["tipo"] == "tutor":
        return redirect(url_for("menu_tutor", r="error"))
    if session["tipo"] == "laboratorista":
        return redirect(url_for("menu_laboratorista", r="error"))
    return redirect(url_for("menu_alumno", r="error"))


# -------------------------------------------------------------------
# Módulo de alumno / externo
# -------------------------------------------------------------------

@app.route("/menu_alumno")
def menu_alumno():
    """Carga el panel del alumno o externo, sus datos y los grupos disponibles cuando aplica."""
    if not login_requerido():
        return redirect(url_for("index"))

    if session["tipo"] not in ["alumno", "externo"]:
        return redirect(url_for("index"))

    foto = obtener_foto_usuario(session["matricula"])
    grupos_disponibles = []
    datos_usuario = usuario.obtener_mis_datos(session["matricula"], session["tipo"])

    if session["tipo"] == "alumno":
        grupos_disponibles = usuario.obtener_grupos_disponibles_para_alumno(
            session["matricula"]
        )

    return render_template(
        "Alm/menu.html",
        nombre=session["nombre"],
        matricula=session["matricula"],
        email=session["email"],
        foto=foto,
        grupos_disponibles=grupos_disponibles,
        datos_usuario=datos_usuario
    )

@app.route("/generar_qr_alumno")
def generar_qr_alumno():
    """Genera la imagen PNG del código QR del usuario autenticado."""
    if not login_requerido():
        return redirect(url_for("index"))

    if session["tipo"] not in ["alumno", "externo"]:
        return redirect(url_for("index"))

    contenido = f'{session["matricula"]}-{session["email"]}'

    qr = qrcode.QRCode(version=1, box_size=10, border=4)
    qr.add_data(contenido)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")

    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    buffer.seek(0)

    return send_file(buffer, mimetype="image/png")

@app.route("/descargar_qr_pdf")
def descargar_qr_pdf():
    """Genera el PDF oficial del QR para alumnos y externos usando la plantilla del proyecto."""
    if not login_requerido():
        return redirect(url_for("index"))

    if session["tipo"] not in ["alumno", "externo"]:
        return redirect(url_for("index"))

    nombre_alumno = session["nombre"]
    matricula = session["matricula"]
    email = session["email"]
    carrera = session["carrera"] if session.get("carrera") else "Externo"

    contenido_qr = f"{matricula}-{email}"

    qr = qrcode.QRCode(version=1, box_size=10, border=4)
    qr.add_data(contenido_qr)
    qr.make(fit=True)

    img_qr = qr.make_image(fill_color="black", back_color="white")
    qr_buffer = io.BytesIO()
    img_qr.save(qr_buffer, format="PNG")
    qr_buffer.seek(0)

    pdf = PdfQrAlumno(
        logo_path="static/resources/logo.png",
        nombre_escuela="Instituto Tecnológico Superior de Atlixco - ITSA",
        nombre_alumno=nombre_alumno,
        matricula=matricula,
        email=email,
        carrera=carrera,
        qr_buffer=qr_buffer,
    )

    pdf_buffer = pdf.construir()

    return send_file(
        pdf_buffer,
        as_attachment=True,
        download_name=f"QR_{matricula}.pdf",
        mimetype="application/pdf",
    )

@app.route("/subir_imagen_alumno", methods=["POST"])
def subir_imagen_alumno():
    """Guarda o reemplaza la foto del alumno o externo dentro de la carpeta de uploads."""
    if not login_requerido():
        return redirect(url_for("index"))

    if session["tipo"] not in ["alumno", "externo"]:
        return redirect(url_for("index"))

    if "foto" not in request.files:
        return redirect(url_for("menu_alumno", r="image_empty"))

    archivo = request.files["foto"]

    if archivo.filename == "":
        return redirect(url_for("menu_alumno", r="image_empty"))

    extension = (
        archivo.filename.rsplit(".", 1)[1].lower() if "." in archivo.filename else ""
    )

    extensiones_validas = ["png", "jpg", "jpeg", "webp"]
    if extension not in extensiones_validas:
        return redirect(url_for("menu_alumno", r="image_wrong"))

    matricula = session["matricula"]
    carpeta = app.config["UPLOAD_FOLDER"]

    for ext in extensiones_validas:
        ruta_vieja = os.path.join(carpeta, f"{matricula}.{ext}")
        if os.path.exists(ruta_vieja):
            os.remove(ruta_vieja)

    nombre_final = f"{matricula}.{extension}"
    ruta_final = os.path.join(carpeta, nombre_final)

    archivo.save(ruta_final)

    return redirect(url_for("menu_alumno", r="image_ok"))

@app.route("/inscribirme_grupo/<int:grupo_id>", methods=["POST"])
def inscribirme_grupo(grupo_id):
    """Permite que un alumno se inscriba a un grupo disponible desde su panel."""
    if not login_requerido("alumno"):
        return redirect(url_for("index"))

    resultado = usuario.inscribir_alumno_a_grupo(session["matricula"], grupo_id)

    if resultado["status"] == "ok":
        return redirect(url_for("menu_alumno", r="group_joined"))

    if resultado["status"] == "student_not_found":
        return redirect(url_for("menu_alumno", r="student_not_found"))

    if resultado["status"] == "group_not_found":
        return redirect(url_for("menu_alumno", r="group_not_found"))

    if resultado["status"] == "wrong_career":
        return redirect(url_for("menu_alumno", r="wrong_career"))

    if resultado["status"] == "exists":
        return redirect(url_for("menu_alumno", r="student_group_exists"))

    return redirect(url_for("menu_alumno", r="error"))


# -------------------------------------------------------------------
# Módulo de tutor
# -------------------------------------------------------------------

@app.route("/menu_tutor")
def menu_tutor():
    """Carga el panel de tutor con sus grupos, catálogos y solicitudes pendientes."""
    if not login_requerido("tutor"):
        return redirect(url_for("index"))

    pagina = request.args.get("page", 1, type=int)
    por_pagina = 10
    offset = (pagina - 1) * por_pagina

    total_solicitudes = usuario.contar_solicitudes_visibles()
    total_paginas = max(1, (total_solicitudes + por_pagina - 1) // por_pagina)

    if pagina < 1:
        pagina = 1
    if pagina > total_paginas:
        pagina = total_paginas
        offset = (pagina - 1) * por_pagina

    solicitudes = usuario.obtener_solicitudes_pendientes(
        limit=por_pagina, offset=offset
    )
    grupos = tutor_manager.obtener_grupos_tutor(session["matricula"])
    carreras = tutor_manager.obtener_carreras_activas()
    materias = tutor_manager.obtener_materias_activas()
    datos_usuario = usuario.obtener_mis_datos(session["matricula"], session["tipo"])
    grupos_sin_tutor = tutor_manager.obtener_grupos_sin_tutor()
    grupos_con_alumnos = []

    for g in grupos:
        alumnos_grupo = tutor_manager.obtener_alumnos_por_grupo(g["id"])
        grupo_temp = dict(g)
        grupo_temp["alumnos"] = alumnos_grupo
        grupos_con_alumnos.append(grupo_temp)

    return render_template(
        "Tutor/menu.html",
        nombre=session["nombre"],
        solicitudes=solicitudes,
        grupos=grupos_con_alumnos,
        carreras=carreras,
        materias=materias,
        datos_usuario=datos_usuario,
        grupos_sin_tutor=grupos_sin_tutor,
        pagina=pagina,
        total_paginas=total_paginas,
    )

@app.route("/buscar_alumno_qr", methods=["POST"])
def buscar_alumno_qr():
    """Procesa la lectura de un QR desde tutor para identificar al alumno o externo."""
    if not login_requerido("tutor"):
        return jsonify({"status": "error", "message": "Acceso no autorizado"}), 403

    usuario.limpiar_externos_expirados()

    data = request.get_json(silent=True)

    if not data:
        return (
            jsonify({"status": "error", "message": "No llegaron datos al servidor"}),
            400,
        )

    contenido_qr = data.get("contenido_qr", "").strip()

    if contenido_qr == "":
        return (
            jsonify({"status": "error", "message": "No se recibió información del QR"}),
            400,
        )

    if "-" not in contenido_qr:
        return jsonify({"status": "error", "message": "Formato de QR inválido"}), 400

    matricula, correo = contenido_qr.split("-", 1)
    matricula = matricula.strip()
    correo = correo.strip()

    if matricula == "":
        return jsonify({"status": "error", "message": "La matrícula está vacía"}), 400

    alumno = usuario.buscar_alumno_por_matricula(matricula)

    if alumno:
        return jsonify(
            {
                "status": "ok",
                "tipo": "alumno",
                "alumno": {
                    "matricula": alumno["matricula"],
                    "nombre": alumno["nombre_completo"],
                    "email": alumno["email"],
                    "carrera": alumno["carrera"],
                },
                "correo_qr": correo,
            }
        )

    externo = usuario.buscar_externo_por_matricula(matricula)

    if externo:
        return jsonify(
            {
                "status": "ok",
                "tipo": "externo",
                "alumno": {
                    "matricula": externo["matricula"],
                    "nombre": externo["nombre_completo"],
                    "email": externo["email"],
                    "carrera": "Externo",
                },
                "correo_qr": correo,
            }
        )

    return (
        jsonify(
            {
                "status": "not_found",
                "message": "No se encontró al estudiante o externo"
            }
        ),
        404,
    )

@app.route("/aprobar_solicitud/<int:solicitud_id>", methods=["POST"])
def aprobar_solicitud(solicitud_id):
    """Aprueba una solicitud de registro pendiente desde el módulo de tutor."""
    if not login_requerido("tutor"):
        return redirect(url_for("index"))

    resultado = usuario.aprobar_solicitud(solicitud_id)

    if resultado["status"] == "ok":
        return redirect(url_for("menu_tutor", r="approved"))

    if resultado["status"] == "not_found":
        return redirect(url_for("menu_tutor", r="not_found"))

    return redirect(url_for("menu_tutor", r="error"))

@app.route("/rechazar_solicitud/<int:solicitud_id>", methods=["POST"])
def rechazar_solicitud(solicitud_id):
    """Rechaza una solicitud de registro pendiente desde el módulo de tutor."""
    if not login_requerido("tutor"):
        return redirect(url_for("index"))

    resultado = usuario.rechazar_solicitud(solicitud_id)

    if resultado["status"] == "ok":
        return redirect(url_for("menu_tutor", r="rejected"))

    return redirect(url_for("menu_tutor", r="error"))

@app.route("/agregar_carrera", methods=["POST"])
def agregar_carrera():
    """Registra una nueva carrera disponible dentro del sistema."""
    if not login_requerido("tutor"):
        return redirect(url_for("index"))

    nombre = request.form.get("nombre", "").strip()

    if nombre == "":
        return redirect(url_for("menu_tutor", r="empty_request"))

    resultado = tutor_manager.agregar_carrera(nombre)

    if resultado["status"] == "ok":
        return redirect(url_for("menu_tutor", r="ok"))

    if resultado["status"] == "exists":
        return redirect(url_for("menu_tutor", r="career_exists"))

    return redirect(url_for("menu_tutor", r="error"))

@app.route("/agregar_materia", methods=["POST"])
def agregar_materia():
    """Crea una nueva materia vinculada a una carrera y semestre."""
    if not login_requerido("tutor"):
        return redirect(url_for("index"))

    carrera_id = request.form.get("carrera_id", "").strip()
    nombre = request.form.get("nombre", "").strip()
    semestre = request.form.get("semestre", "").strip()

    if carrera_id == "" or nombre == "" or semestre == "":
        return redirect(url_for("menu_tutor", r="empty_request"))

    resultado = tutor_manager.agregar_materia(int(carrera_id), nombre, semestre)

    if resultado["status"] == "ok":
        return redirect(url_for("menu_tutor", r="ok"))

    if resultado["status"] == "exists":
        return redirect(url_for("menu_tutor", r="subject_exists"))

    if resultado["status"] == "career_not_found":
        return redirect(url_for("menu_tutor", r="not_found"))

    return redirect(url_for("menu_tutor", r="error"))

@app.route("/crear_grupo", methods=["POST"])
def crear_grupo():
    """Crea un grupo académico con carrera, materia, turno y periodo."""
    if not login_requerido("tutor"):
        return redirect(url_for("index"))

    carrera_id = request.form.get("carrera_id", "").strip()
    materia_id = request.form.get("materia_id", "").strip()
    semestre = request.form.get("semestre", "").strip()
    grupo = request.form.get("grupo", "").strip()
    turno = request.form.get("turno", "").strip()
    periodo = request.form.get("periodo", "").strip()

    if (
        carrera_id == ""
        or materia_id == ""
        or semestre == ""
        or grupo == ""
        or turno == ""
        or periodo == ""
    ):
        return redirect(url_for("menu_tutor", r="empty_request"))

    datos = {
        "carrera_id": int(carrera_id),
        "materia_id": int(materia_id),
        "semestre": semestre,
        "grupo": grupo,
        "turno": turno,
        "periodo": periodo,
    }

    resultado = tutor_manager.crear_grupo(datos)

    if resultado["status"] == "ok":
        return redirect(url_for("menu_tutor", r="ok"))

    if resultado["status"] == "exists":
        return redirect(url_for("menu_tutor", r="group_exists"))

    return redirect(url_for("menu_tutor", r="error"))

@app.route("/asignarme_grupo/<int:grupo_id>", methods=["POST"])
def asignarme_grupo(grupo_id):
    """Asigna al tutor autenticado a un grupo disponible."""
    if not login_requerido("tutor"):
        return redirect(url_for("index"))

    resultado = tutor_manager.asignar_tutor_a_grupo(grupo_id, session["matricula"])

    if resultado["status"] == "ok":
        return redirect(url_for("menu_tutor", r="group_assigned"))

    if resultado["status"] == "not_found":
        return redirect(url_for("menu_tutor", r="not_found"))

    if resultado["status"] == "occupied":
        return redirect(url_for("menu_tutor", r="group_occupied"))

    return redirect(url_for("menu_tutor", r="error"))

@app.route("/agregar_alumno_grupo", methods=["POST"])
def agregar_alumno_grupo():
    """Agrega manualmente a un alumno dentro de un grupo."""
    if not login_requerido("tutor"):
        return redirect(url_for("index"))

    alumno_matricula = request.form.get("alumno_matricula", "").strip()
    grupo_id = request.form.get("grupo_id", "").strip()

    if alumno_matricula == "" or grupo_id == "":
        return redirect(url_for("menu_tutor", r="empty_request"))

    resultado = tutor_manager.agregar_alumno_a_grupo(alumno_matricula, int(grupo_id))

    if resultado["status"] == "ok":
        return redirect(url_for("menu_tutor", r="student_group_ok"))

    if resultado["status"] == "student_not_found":
        return redirect(url_for("menu_tutor", r="student_not_found"))

    if resultado["status"] == "group_not_found":
        return redirect(url_for("menu_tutor", r="group_not_found"))

    if resultado["status"] == "exists":
        return redirect(url_for("menu_tutor", r="student_group_exists"))

    return redirect(url_for("menu_tutor", r="error"))

@app.route("/quitar_alumno_grupo", methods=["POST"])
def quitar_alumno_grupo():
    """Retira a un alumno de un grupo específico."""
    if not login_requerido("tutor"):
        return redirect(url_for("index"))

    alumno_matricula = request.form.get("alumno_matricula", "").strip()
    grupo_id = request.form.get("grupo_id", "").strip()

    if alumno_matricula == "" or grupo_id == "":
        return redirect(url_for("menu_tutor", r="empty_request"))

    resultado = tutor_manager.quitar_alumno_de_grupo(alumno_matricula, int(grupo_id))

    if resultado["status"] == "ok":
        return redirect(url_for("menu_tutor", r="ok"))

    return redirect(url_for("menu_tutor", r="error"))

@app.route("/registrar_personal", methods=["POST"])
def registrar_personal():
    """Registra directamente personal interno como tutor o laboratorista."""
    if not login_requerido("tutor"):
        return redirect(url_for("index"))

    nombre = request.form.get("nombre", "").strip()
    ap_paterno = request.form.get("ap_paterno", "").strip()
    ap_materno = request.form.get("ap_materno", "").strip()
    correo_temporal = request.form.get("correo_temporal", "").strip()
    password = request.form.get("password", "").strip()
    tipo = request.form.get("tipo", "").strip()

    if (
        nombre == ""
        or ap_paterno == ""
        or ap_materno == ""
        or correo_temporal == ""
        or password == ""
        or tipo == ""
    ):
        return redirect(url_for("menu_tutor", r="empty_request"))

    datos = {
        "nombre": nombre,
        "ap_paterno": ap_paterno,
        "ap_materno": ap_materno,
        "correo_temporal": correo_temporal,
        "password": password,
        "tipo": tipo,
    }

    resultado = usuario.registrar_personal_directo(datos)

    if resultado["status"] == "ok":
        return redirect(url_for("menu_tutor", r="ok"))

    return redirect(url_for("menu_tutor", r="error"))

@app.route("/tutor/seguimiento/<int:grupo_id>/<parcial>")
def abrir_seguimiento(grupo_id, parcial):
    """Abre o crea el seguimiento académico de un grupo para el parcial indicado."""
    if not login_requerido("tutor"):
        return redirect(url_for("index"))

    seguimiento = tutor_manager.obtener_o_crear_seguimiento(
        grupo_id, session["matricula"], parcial
    )

    alumnos = tutor_manager.obtener_alumnos_de_grupo(grupo_id)
    detalle = tutor_manager.obtener_detalle_seguimiento(seguimiento["id"])

    detalle_map = {}
    for d in detalle:
        detalle_map[d["alumno_matricula"]] = d

    return render_template(
        "Tutor/seguimiento.html",
        seguimiento=seguimiento,
        alumnos=alumnos,
        detalle_map=detalle_map,
    )

@app.route("/guardar_seguimiento", methods=["POST"])
def guardar_seguimiento():
    """Guarda el detalle de seguimiento académico capturado por el tutor."""
    if not login_requerido("tutor"):
        return redirect(url_for("index"))

    seguimiento_id = request.form.get("seguimiento_id", "").strip()
    observaciones = request.form.get("observaciones", "").strip()
    alumnos_matricula = request.form.getlist("alumno_matricula")

    if seguimiento_id == "" or not alumnos_matricula:
        return redirect(url_for("menu_tutor", r="error"))

    alumnos = []

    for matricula in alumnos_matricula:
        alumno = {
            "alumno_matricula": matricula,
            "asistencia": (
                1 if request.form.get(f"asistencia_{matricula}") == "1" else 0
            ),
            "canalizacion_tutoria_individual": (
                1 if request.form.get(f"canalizacion_{matricula}") == "1" else 0
            ),
            "problemas_economicos": (
                1 if request.form.get(f"prob_economicos_{matricula}") == "1" else 0
            ),
            "problemas_salud": (
                1 if request.form.get(f"prob_salud_{matricula}") == "1" else 0
            ),
            "problemas_emocionales": (
                1 if request.form.get(f"prob_emocionales_{matricula}") == "1" else 0
            ),
            "problemas_familiares": (
                1 if request.form.get(f"prob_familiares_{matricula}") == "1" else 0
            ),
            "problemas_academicos": (
                1 if request.form.get(f"prob_academicos_{matricula}") == "1" else 0
            ),
            "beca_jcf": 1 if request.form.get(f"beca_jcf_{matricula}") == "1" else 0,
            "beca_manutencion": (
                1 if request.form.get(f"beca_manutencion_{matricula}") == "1" else 0
            ),
            "beca_jmf": 1 if request.form.get(f"beca_jmf_{matricula}") == "1" else 0,
            "beca_otro": 1 if request.form.get(f"beca_otro_{matricula}") == "1" else 0,
            "apoyo_psicologico": (
                1 if request.form.get(f"apoyo_psicologico_{matricula}") == "1" else 0
            ),
            "apoyo_medico": (
                1 if request.form.get(f"apoyo_medico_{matricula}") == "1" else 0
            ),
            "apoyo_otro": (
                1 if request.form.get(f"apoyo_otro_{matricula}") == "1" else 0
            ),
            "promedio": request.form.get(f"promedio_{matricula}") or None,
            "materias_reprobadas": int(
                request.form.get(f"reprobadas_{matricula}") or 0
            ),
            "estatus_aprobado": (
                1 if request.form.get(f"aprobado_{matricula}") == "1" else 0
            ),
            "posible_desercion": (
                1 if request.form.get(f"desercion_{matricula}") == "1" else 0
            ),
            "comentario": request.form.get(f"comentario_{matricula}", "").strip(),
        }

        alumnos.append(alumno)

    resultado = tutor_manager.guardar_detalle_seguimiento(
        int(seguimiento_id), alumnos, observaciones
    )

    if resultado["status"] == "ok":
        return redirect(url_for("menu_tutor", r="ok"))

    return redirect(url_for("menu_tutor", r="error"))

@app.route("/tutor/seguimiento/<int:seguimiento_id>/exportar-pdf", methods=["POST"])
def exportar_seguimiento_pdf(seguimiento_id):
    if not login_requerido("tutor"):
        return redirect(url_for("index"))
    

    seguimiento, alumnos, detalle_map = construir_datos_exportacion_desde_formulario()

    pdf = PdfSeguimientoAcademico(
        "static/img/logo_itsa.png",
        "Instituto Tecnológico Superior de Atlixco",
        seguimiento,
        alumnos,
        detalle_map
    )

    buffer = pdf.construir()

    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"seguimiento_{seguimiento_id}.pdf",
        mimetype="application/pdf"
    )

@app.route("/tutor/seguimiento/<int:seguimiento_id>/exportar-excel", methods=["POST"])
def exportar_seguimiento_excel(seguimiento_id):
    if not login_requerido("tutor"):
        return redirect(url_for("index"))
   
    seguimiento, alumnos, detalle_map = construir_datos_exportacion_desde_formulario()

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from io import BytesIO

    wb = Workbook()
    ws = wb.active
    ws.title = "Seguimiento"

    fill_rojo = PatternFill(fill_type="solid", fgColor="B30000")
    font_blanco = Font(color="FFFFFF", bold=True)
    font_negrita = Font(bold=True)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    ws["A1"] = "Seguimiento Académico"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A2"] = "Grupo ID"
    ws["B2"] = seguimiento.get("grupo_id", "")
    ws["C2"] = "Parcial"
    ws["D2"] = seguimiento.get("parcial", "")
    ws["E2"] = "Total estudiantes"
    ws["F2"] = len(alumnos)

    encabezados = [
        "Matrícula", "Nombre", "Asistencia", "Promedio", "Reprobadas",
        "Aprobado", "Posible deserción", "Problemas", "Becas", "Apoyos", "Comentario"
    ]

    fila_inicio = 5
    for col, encabezado in enumerate(encabezados, 1):
        celda = ws.cell(row=fila_inicio, column=col, value=encabezado)
        celda.fill = fill_rojo
        celda.font = font_blanco
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = border

    def compactar_problemas(d):
        problemas = []
        if d.get("problemas_economicos") == 1:
            problemas.append("Económicos")
        if d.get("problemas_salud") == 1:
            problemas.append("Salud")
        if d.get("problemas_emocionales") == 1:
            problemas.append("Emocionales")
        if d.get("problemas_familiares") == 1:
            problemas.append("Familiares")
        if d.get("problemas_academicos") == 1:
            problemas.append("Académicos")
        if not problemas:
            return "Ninguno"
        if len(problemas) == 1:
            return problemas[0]
        if len(problemas) == 2:
            return f"{problemas[0]} y {problemas[1]}"
        return f"Varios ({len(problemas)}): " + ", ".join(problemas)

    def compactar_becas(d):
        becas = []
        if d.get("beca_jcf") == 1:
            becas.append("JCF")
        if d.get("beca_manutencion") == 1:
            becas.append("Manutención")
        if d.get("beca_jmf") == 1:
            becas.append("JMF")
        if d.get("beca_otro") == 1:
            becas.append("Otra")
        return ", ".join(becas) if becas else "Ninguna"

    def compactar_apoyos(d):
        apoyos = []
        if d.get("apoyo_psicologico") == 1:
            apoyos.append("Psicológico")
        if d.get("apoyo_medico") == 1:
            apoyos.append("Médico")
        if d.get("apoyo_otro") == 1:
            apoyos.append("Otro")
        return ", ".join(apoyos) if apoyos else "Ninguno"

    fila = fila_inicio + 1
    for alumno in alumnos:
        d = detalle_map.get(alumno["matricula"], {})

        ws.cell(row=fila, column=1, value=alumno["matricula"])
        ws.cell(row=fila, column=2, value=alumno["nombre_completo"])
        ws.cell(row=fila, column=3, value="Sí" if d.get("asistencia") == 1 else "No")
        ws.cell(row=fila, column=4, value=d.get("promedio", ""))
        ws.cell(row=fila, column=5, value=d.get("materias_reprobadas", 0))
        ws.cell(row=fila, column=6, value="Sí" if d.get("estatus_aprobado") == 1 else "No")
        ws.cell(row=fila, column=7, value="Sí" if d.get("posible_desercion") == 1 else "No")
        ws.cell(row=fila, column=8, value=compactar_problemas(d))
        ws.cell(row=fila, column=9, value=compactar_becas(d))
        ws.cell(row=fila, column=10, value=compactar_apoyos(d))
        ws.cell(row=fila, column=11, value=d.get("comentario", ""))

        for col in range(1, 12):
            celda = ws.cell(row=fila, column=col)
            celda.border = border
            celda.alignment = Alignment(vertical="top", wrap_text=True)

        fila += 1

    fila += 1
    ws.cell(row=fila, column=1, value="Observaciones generales").font = font_negrita
    ws.cell(row=fila + 1, column=1, value=seguimiento.get("observaciones", ""))

    anchos = {
        "A": 18, "B": 30, "C": 12, "D": 10, "E": 12,
        "F": 12, "G": 18, "H": 30, "I": 20, "J": 20, "K": 35
    }

    for col, ancho in anchos.items():
        ws.column_dimensions[col].width = ancho

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=f"seguimiento_{seguimiento_id}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

# -------------------------------------------------------------------
# Módulo de laboratorista
# -------------------------------------------------------------------

@app.route("/menu_laboratorista")
def menu_laboratorista():
    """Carga el panel del laboratorista, actualiza préstamos vencidos y arma la paginación de equipos."""
    if not login_requerido("laboratorista"):
        return redirect(url_for("index"))

    vencidos = laboratorista_manager.actualizar_estado_prestamos()

    pagina = request.args.get("page", 1, type=int)
    por_pagina = 25

    todos_los_equipos = laboratorista_manager.obtener_equipos()
    total_equipos = len(todos_los_equipos)
    total_paginas = max(1, (total_equipos + por_pagina - 1) // por_pagina)

    if pagina < 1:
        pagina = 1
    if pagina > total_paginas:
        pagina = total_paginas

    inicio = (pagina - 1) * por_pagina
    fin = inicio + por_pagina

    equipos = todos_los_equipos[inicio:fin]
    equipos_disponibles = laboratorista_manager.obtener_equipos_disponibles()
    datos_usuario = usuario.obtener_mis_datos(session["matricula"], session["tipo"])

    return render_template(
        "Laboratorista/menu.html",
        nombre=session["nombre"],
        equipos=equipos,
        equipos_disponibles=equipos_disponibles,
        pagina=pagina,
        total_paginas=total_paginas,
        dispositivos_vencidos=vencidos,
        datos_usuario=datos_usuario
    )

@app.route("/laboratorista/registrar_equipo", methods=["POST"])
def registrar_equipo_laboratorio():
    """Registra un nuevo equipo de laboratorio con sus datos principales."""
    if not login_requerido("laboratorista"):
        return redirect(url_for("index"))

    codigo_equipo = request.form.get("codigo_equipo", "").strip()
    nombre_equipo = request.form.get("nombre_equipo", "").strip()
    categoria = request.form.get("categoria", "").strip()
    marca = request.form.get("marca", "").strip()
    modelo = request.form.get("modelo", "").strip()
    numero_serie = request.form.get("numero_serie", "").strip()
    descripcion = request.form.get("descripcion", "").strip()
    estado = request.form.get("estado", "disponible").strip()
    ubicacion = request.form.get("ubicacion", "").strip()

    if codigo_equipo == "" or nombre_equipo == "" or categoria == "":
        return redirect(url_for("menu_laboratorista", r="empty_request"))

    datos = {
        "codigo_equipo": codigo_equipo,
        "nombre_equipo": nombre_equipo,
        "categoria": categoria,
        "marca": marca if marca != "" else None,
        "modelo": modelo if modelo != "" else None,
        "numero_serie": numero_serie if numero_serie != "" else None,
        "descripcion": descripcion if descripcion != "" else None,
        "estado": estado if estado != "" else "disponible",
        "ubicacion": ubicacion if ubicacion != "" else None,
    }

    resultado = laboratorista_manager.registrar_equipo(datos)

    if resultado["status"] == "ok":
        return redirect(url_for("menu_laboratorista", r="ok"))

    if resultado["status"] == "exists":
        return redirect(url_for("menu_laboratorista", r="equipment_exists"))

    return redirect(url_for("menu_laboratorista", r="error"))

@app.route("/laboratorista/buscar_usuario_qr", methods=["POST"])
def buscar_usuario_qr_laboratorista():
    """Consulta al usuario leído desde un QR para preparar el préstamo de un equipo."""
    if not login_requerido("laboratorista"):
        return jsonify({"status": "error", "message": "Acceso no autorizado"}), 403

    data = request.get_json(silent=True)

    if not data:
        return jsonify({"status": "error", "message": "No llegaron datos"}), 400

    contenido_qr = data.get("contenido_qr", "").strip()

    if contenido_qr == "":
        return jsonify({"status": "error", "message": "QR vacío"}), 400

    if "-" not in contenido_qr:
        return jsonify({"status": "error", "message": "Formato de QR inválido"}), 400

    matricula, correo = contenido_qr.split("-", 1)
    matricula = matricula.strip()

    usuario_qr = laboratorista_manager.buscar_usuario_por_matricula(matricula)

    if not usuario_qr:
        return jsonify({"status": "not_found", "message": "Usuario no encontrado"}), 404

    return jsonify({"status": "ok", "usuario": usuario_qr})

@app.route("/laboratorista/registrar_prestamo", methods=["POST"])
def registrar_prestamo_laboratorista():
    """Registra un préstamo de equipo a partir del usuario y dispositivo seleccionados."""
    if not login_requerido("laboratorista"):
        return redirect(url_for("index"))

    equipo_id = request.form.get("equipo_id", "").strip()
    matricula_usuario = request.form.get("matricula_usuario", "").strip()
    tipo_usuario = request.form.get("tipo_usuario", "").strip()
    fecha_limite = request.form.get("fecha_limite", "").strip()
    observaciones = request.form.get("observaciones", "").strip()

    if equipo_id == "" or matricula_usuario == "" or tipo_usuario == "":
        return redirect(url_for("menu_laboratorista", r="empty_request"))

    resultado = laboratorista_manager.registrar_prestamo(
        int(equipo_id), matricula_usuario, tipo_usuario, fecha_limite, observaciones
    )

    if resultado["status"] == "ok":
        return redirect(url_for("menu_laboratorista", r="loan_ok"))

    if resultado["status"] == "team_not_found":
        return redirect(url_for("menu_laboratorista", r="group_not_found"))

    if resultado["status"] in ["not_available", "already_loaned"]:
        return redirect(url_for("menu_laboratorista", r="equipment_busy"))

    return redirect(url_for("menu_laboratorista", r="error"))

@app.route("/laboratorista/equipos_estado")
def equipos_estado_laboratorista():
    """Devuelve por JSON el estado actualizado de los equipos y los préstamos vencidos."""
    if not login_requerido("laboratorista"):
        return jsonify({"status": "error"}), 403

    vencidos = laboratorista_manager.actualizar_estado_prestamos()
    equipos = laboratorista_manager.obtener_equipos()

    return jsonify({"status": "ok", "equipos": equipos, "vencidos": vencidos})

@app.route("/laboratorista/generar_excel")
def generar_excel_laboratorista():
    """Genera y descarga la bitácora de laboratorio en Excel a partir de la plantilla base."""
    if not login_requerido("laboratorista"):
        return redirect(url_for("index"))

    template_path = os.path.join(
        "static", "resources", "Bitácora de uso general (Blanco).xlsx"
    )
    output_path = os.path.join(
        "static", "resources", f"bitacora_generada_{session['matricula']}.xlsx"
    )

    archivo = laboratorista_manager.generar_excel_bitacora(
        nombre_laboratorio="Laboratorio de Cómputo",
        periodo=datetime.now().strftime("%Y"),
        responsable=session["nombre"],
        template_path=template_path,
        output_path=output_path,
    )

    return send_file(
        archivo,
        as_attachment=True,
        download_name="bitacora_laboratorio.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=3500, debug=True)
